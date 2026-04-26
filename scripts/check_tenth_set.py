from docx import Document

# 读取第十套题目
doc = Document('c:\\Users\\86183\\Desktop\\111\\第十套.docx')

print("=== 文档内容 ===")

# 打印前50个段落，了解文档结构
for i, para in enumerate(doc.paragraphs[:50]):
    if para.text.strip():
        print(f"[{i}] {para.text}")

print("\n=== 文档表格 ===")

# 检查是否有表格
for table_idx, table in enumerate(doc.tables):
    print(f"表格 {table_idx}:")
    # 打印前5行
    for row_idx, row in enumerate(table.rows[:5]):
        row_data = [cell.text.strip() for cell in row.cells]
        print(f"  行 {row_idx}: {row_data}")
