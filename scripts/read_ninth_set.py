
from docx import Document
import json
import os

def read_ninth_set():
    print('=' * 60)
    print('读取第九套题库')
    print('=' * 60)
    print()
    
    # 读取第九套题库
    doc_path = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), '第九套.docx')
    doc = Document(doc_path)
    
    print("=== 文档段落内容 ===\n")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            print(f"[{i}] {text}")
    
    print("\n\n=== 文档表格内容 ===\n")
    for table_idx, table in enumerate(doc.tables):
        print(f"\n--- 表格 {table_idx} ---")
        for row_idx, row in enumerate(table.rows):
            row_data = [cell.text.strip() for cell in row.cells]
            print(f"行 {row_idx}: {row_data}")
    
    print('=' * 60)

if __name__ == '__main__':
    read_ninth_set()
